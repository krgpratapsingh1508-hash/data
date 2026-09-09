import streamlit as st
import pandas as pd
import sqlite3
import json

st.set_page_config(page_title="NEP Master Data System", page_icon="🎓", layout="wide")

# =========================================================================
# 🎨 प्रोफेशनल UI स्टाइलिंग (पूरे ऐप में कस्टम CSS)
# =========================================================================
st.markdown("""
<style>
    /* मुख्य कंटेनर पैडिंग */
    .block-container {
        padding-top: 1.5rem;
        padding-bottom: 3rem;
    }
    /* हैडिंग्स */
    h1, h2, h3 {
        font-family: 'Segoe UI', 'Trebuchet MS', sans-serif;
        letter-spacing: -0.3px;
    }
    h1 { color: #1a3c6e; }
    /* बटन */
    div.stButton > button {
        border-radius: 8px;
        font-weight: 600;
        border: 1px solid #d0d7e2;
        transition: all 0.15s ease-in-out;
    }
    div.stButton > button:hover {
        border-color: #1a73e8;
        color: #1a73e8;
    }
    div.stButton > button[kind="primary"] {
        background-color: #1a73e8;
    }
    /* डाउनलोड बटन */
    div.stDownloadButton > button {
        border-radius: 8px;
        font-weight: 600;
        background-color: #0f9d58;
        color: white;
        border: none;
    }
    div.stDownloadButton > button:hover {
        background-color: #0c8043;
        color: white;
    }
    /* साइडबार */
    section[data-testid="stSidebar"] {
        background-color: #f6f8fb;
        border-right: 1px solid #e3e8ef;
    }
    /* डेटाफ़्रेम / टेबल कार्ड जैसा दिखे */
    div[data-testid="stDataFrame"] {
        border: 1px solid #e3e8ef;
        border-radius: 10px;
        overflow: hidden;
    }
    /* मेट्रिक कार्ड्स */
    div[data-testid="stMetric"] {
        background-color: #f8fafc;
        border: 1px solid #e3e8ef;
        border-radius: 10px;
        padding: 12px 16px;
    }
    /* एक्सपैंडर */
    div[data-testid="stExpander"] {
        border: 1px solid #e3e8ef;
        border-radius: 10px;
    }
    /* टैब्स */
    button[data-baseweb="tab"] {
        font-weight: 600;
    }
    /* फुटर क्रेडिट */
    .app-footer {
        text-align: center;
        color: #8a94a6;
        font-size: 12.5px;
        padding: 18px 0 4px 0;
        border-top: 1px solid #e3e8ef;
        margin-top: 30px;
    }

    /* ============== 🔐 लॉगिन स्क्रीन स्टाइलिंग ============== */
    @keyframes floatIn {
        0% { opacity: 0; transform: translateY(14px); }
        100% { opacity: 1; transform: translateY(0); }
    }
    @keyframes gradientShift {
        0% { background-position: 0% 50%; }
        50% { background-position: 100% 50%; }
        100% { background-position: 0% 50%; }
    }
    .login-hero {
        text-align: center;
        animation: floatIn 0.6s ease-out;
        margin-bottom: 6px;
    }
    .login-hero .emoji-badge {
        font-size: 46px;
        display: inline-block;
        animation: floatIn 0.5s ease-out;
    }
    .login-hero .login-logo-img {
        max-height: 92px;
        max-width: 260px;
        border-radius: 14px;
        box-shadow: 0 6px 20px rgba(26, 60, 110, 0.18);
        animation: floatIn 0.5s ease-out;
        object-fit: contain;
        background: #ffffff;
        padding: 6px;
    }
    .login-hero h1 {
        font-size: 34px;
        font-weight: 800;
        margin: 6px 0 2px 0;
        background: linear-gradient(270deg, #1a73e8, #6a4cff, #0f9d58, #1a73e8);
        background-size: 600% 600%;
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        background-clip: text;
        animation: gradientShift 6s ease infinite;
    }
    .login-hero p {
        color: #6b7688;
        font-size: 15px;
        margin-top: 0;
    }
    div[data-testid="stVerticalBlockBorderWrapper"] {
        animation: floatIn 0.7s ease-out;
    }
    /* लॉगिन कार्ड (st.container(border=True)) को थोड़ा प्रीमियम लुक */
    div[data-testid="stForm"], div[data-testid="stVerticalBlockBorderWrapper"] > div {
        border-radius: 16px !important;
    }
    .login-badge-row {
        display: flex;
        justify-content: center;
        gap: 8px;
        flex-wrap: wrap;
        margin: 10px 0 18px 0;
    }
    .login-badge {
        background: linear-gradient(135deg, #eef3ff, #f3eefc);
        border: 1px solid #dde4f5;
        color: #4a5b8c;
        font-size: 12px;
        font-weight: 600;
        padding: 5px 12px;
        border-radius: 20px;
    }
    .farewell-box {
        text-align: center;
        animation: floatIn 0.5s ease-out;
        background: linear-gradient(135deg, #eafff0, #eef8ff);
        border: 1px solid #cdeedd;
        border-radius: 14px;
        padding: 14px;
        margin-bottom: 18px;
        color: #1a6e3c;
        font-weight: 600;
    }
</style>
""", unsafe_allow_html=True)

def render_footer():
    st.markdown(
        "<div class='app-footer'>🛠️ Professionally Developed &amp; Maintained &nbsp;|&nbsp; "
        "NEP Master Data System © 2026</div>",
        unsafe_allow_html=True
    )

# =========================================================================
# डेटाबेस सेटअप - टेबल्स संरचना (Raw, Permanent और Rules Lock)
# =========================================================================
conn = sqlite3.connect("nep_master_perma_db.db", check_same_thread=False)
cursor = conn.cursor()

# 1. अस्थायी स्टेजिंग स्टोरेज (Panel 1 से Upload होकर यहाँ आएगा)
cursor.execute("""
    CREATE TABLE IF NOT EXISTS raw_store (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        data_json TEXT
    )
""")

# 2. परमानेंट स्टोरेज (Panel 2 से Approve होकर UG/PG यहाँ आएगा)
cursor.execute("""
    CREATE TABLE IF NOT EXISTS perma_store (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        data_json TEXT,
        course_type TEXT
    )
""")

# 3. परमानेंट नियम लॉकिंग स्टोरेज (पुरानी खराब टेबल को डिलीट करके नया बनाने का ऑटो-सिस्टम)
try:
    # चेक करना कि क्या टेबल सही है
    cursor.execute("SELECT panel_prefix FROM locked_rules LIMIT 1")
except sqlite3.OperationalError:
    # अगर कोई भी गड़बड़ (जैसे कॉलम गायब होना) मिले, तो पुरानी टेबल हटा दें
    cursor.execute("DROP TABLE IF EXISTS locked_rules")
    conn.commit()

# अब बिल्कुल सही और नए स्ट्रक्चर के साथ टेबल बनाएं
cursor.execute("""
    CREATE TABLE IF NOT EXISTS locked_rules (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        panel_prefix TEXT UNIQUE,
        rules_json TEXT
    )
""")
conn.commit()

# सुनिश्चित करें कि टेबल बनने के बाद डेटाबेस में बदलाव सुरक्षित (Commit) हो जाएं
conn.commit()

# 4. पैनल-वाइज पासवर्ड + हाइड/अनहाइड स्टोरेज (6 पैनल्स के लिए)
cursor.execute("""
    CREATE TABLE IF NOT EXISTS panel_auth (
        panel_key TEXT PRIMARY KEY,
        password TEXT,
        hidden INTEGER DEFAULT 0
    )
""")
conn.commit()

# डिफ़ॉल्ट पासवर्ड (सिर्फ पहली बार, जब टेबल में डेटा न हो, तभी डाले जाएंगे)
_default_panel_passwords = {
    "p1": "op",       # पहले Operator का पासवर्ड
    "p2": "p2pass",
    "p3": "ug",       # पहले Teacher_UG का पासवर्ड
    "p4": "pg",       # पहले Teacher_PG का पासवर्ड
    "p5": "p5pass",
    "p6": "psv123",   # पहले Admin का पासवर्ड
}
for _pk, _pw in _default_panel_passwords.items():
    cursor.execute("INSERT OR IGNORE INTO panel_auth (panel_key, password, hidden) VALUES (?, ?, 0)", (_pk, _pw))
conn.commit()

# पैनल-की और उसके डिस्प्ले नाम की मैपिंग
PANEL_KEY_TO_NAME = {
    "p1": "📥 1. Entry / Upload Panel",
    "p2": "💻 2. Work / Approve Panel",
    "p3": "🎓 3. UG Panel",
    "p4": "📜 4. PG Panel",
    "p5": "📊 5. Dashboard / Counter Panel",
    "p6": "⚙️ 6. Admin Panel",
}
PANEL_NAME_TO_KEY = {v: k for k, v in PANEL_KEY_TO_NAME.items()}

def get_panel_password(panel_key):
    cursor.execute("SELECT password FROM panel_auth WHERE panel_key = ?", (panel_key,))
    row = cursor.fetchone()
    return row[0] if row else None

def is_panel_hidden(panel_key):
    cursor.execute("SELECT hidden FROM panel_auth WHERE panel_key = ?", (panel_key,))
    row = cursor.fetchone()
    return bool(row[0]) if row else False

def set_panel_password(panel_key, new_password):
    cursor.execute("UPDATE panel_auth SET password = ? WHERE panel_key = ?", (new_password, panel_key))
    conn.commit()

def set_panel_hidden(panel_key, hidden_flag):
    cursor.execute("UPDATE panel_auth SET hidden = ? WHERE panel_key = ?", (1 if hidden_flag else 0, panel_key))
    conn.commit()

# 5. ऐप सेटिंग्स स्टोरेज (Login टाइटल + लोगो — सिर्फ Admin बदल सकता है)
cursor.execute("""
    CREATE TABLE IF NOT EXISTS app_settings (
        setting_key TEXT PRIMARY KEY,
        setting_value TEXT
    )
""")
conn.commit()

def get_app_setting(key, default=None):
    cursor.execute("SELECT setting_value FROM app_settings WHERE setting_key = ?", (key,))
    row = cursor.fetchone()
    return row[0] if row and row[0] is not None else default

def set_app_setting(key, value):
    cursor.execute("""
        INSERT INTO app_settings (setting_key, setting_value) VALUES (?, ?)
        ON CONFLICT(setting_key) DO UPDATE SET setting_value = excluded.setting_value
    """, (key, value))
    conn.commit()

def delete_app_setting(key):
    cursor.execute("DELETE FROM app_settings WHERE setting_key = ?", (key,))
    conn.commit()

# =========================================================================
# परमानेंट डेटा लोड करने का फंक्शन (perma_store से UG/PG डेटा पढ़ने के लिए)
# =========================================================================
def load_permanent_data(course_type):
    cursor.execute("SELECT data_json FROM perma_store WHERE course_type = ?", (course_type,))
    rows = cursor.fetchall()
    if not rows:
        return pd.DataFrame()
    all_records = []
    for (data_json,) in rows:
        try:
            records = json.loads(data_json)
            if isinstance(records, list):
                all_records.extend(records)
            else:
                all_records.append(records)
        except (json.JSONDecodeError, TypeError):
            continue
    if not all_records:
        return pd.DataFrame()
    return pd.DataFrame(all_records)

def load_raw_data():
    cursor.execute("SELECT data_json FROM raw_store ORDER BY id DESC LIMIT 1")
    row = cursor.fetchone()
    if not row or not row[0]:
        return pd.DataFrame()
    try:
        records = json.loads(row[0])
    except (json.JSONDecodeError, TypeError):
        return pd.DataFrame()
    if isinstance(records, dict):
        records = [records]
    if not records:
        return pd.DataFrame()
    return pd.DataFrame(records)

# Session States Management
if "ok" not in st.session_state: st.session_state["ok"] = False
if "deleted_cols" not in st.session_state: st.session_state["deleted_cols"] = []

# --- LOGIN SYSTEM (पैनल-वाइज: हर पैनल का अपना पासवर्ड) ---
# 🔧 फिक्स: पुराने सेशन (जिसमें "ok"=True था लेकिन "panel" key नहीं थी) की वजह से
# KeyError न आए, इसलिए दोनों चीज़ें एक साथ चेक कर रहे हैं
if not st.session_state["ok"] or "panel" not in st.session_state:
    st.session_state["ok"] = False

    left, mid, right = st.columns([1, 1.3, 1])
    with mid:
        # 👋 अगर अभी-अभी Logout किया है, तो एक प्यारा-सा फेयरवेल मैसेज दिखाएं
        if st.session_state.pop("show_farewell", False):
            st.markdown(
                "<div class='farewell-box'>👋 सफलतापूर्वक Logout हो गए! फिर मिलते हैं 😊</div>",
                unsafe_allow_html=True
            )

        # 🎨 Admin द्वारा सेट किया गया टाइटल और लोगो (डिफ़ॉल्ट: इमोजी + "NEP Master Data System")
        _login_title = get_app_setting("login_title", "NEP Master Data System")
        _login_subtitle = get_app_setting("login_subtitle", "अपना पैनल चुनें और आगे बढ़ने के लिए पासवर्ड डालें")
        _logo_b64 = get_app_setting("login_logo_b64")
        _logo_mime = get_app_setting("login_logo_mime", "image/png")

        if _logo_b64:
            badge_html = f'<img src="data:{_logo_mime};base64,{_logo_b64}" class="login-logo-img" />'
        else:
            badge_html = '<div class="emoji-badge">🎓🔒</div>'

        st.markdown(
            f"""
            <div class="login-hero">
                {badge_html}
                <h1>{_login_title}</h1>
                <p>{_login_subtitle}</p>
            </div>
            """,
            unsafe_allow_html=True
        )
        st.markdown(
            """
            <div class="login-badge-row">
                <span class="login-badge">📥 Entry</span>
                <span class="login-badge">💻 Approve</span>
                <span class="login-badge">🎓 UG</span>
                <span class="login-badge">📜 PG</span>
                <span class="login-badge">📊 Dashboard</span>
                <span class="login-badge">⚙️ Admin</span>
            </div>
            """,
            unsafe_allow_html=True
        )

        with st.container(border=True):
            panel_choice_label = st.selectbox("🗂️ पैनल चुनें:", ["-- चुनें --"] + list(PANEL_KEY_TO_NAME.values()))
            pas = st.text_input("🔑 Password:", type="password", placeholder="अपना पासवर्ड यहाँ डालें")
            login_clicked = st.button("🚀 Login करें", use_container_width=True, type="primary")

            if login_clicked:
                if panel_choice_label == "-- चुनें --":
                    st.error("⚠️ कृपया पहले एक पैनल चुनें।")
                else:
                    selected_key = PANEL_NAME_TO_KEY[panel_choice_label]
                    correct_pw = get_panel_password(selected_key)
                    if correct_pw is not None and pas == correct_pw:
                        st.session_state["ok"] = True
                        st.session_state["panel_key"] = selected_key
                        st.session_state["panel"] = panel_choice_label
                        st.success(f"🎉 स्वागत है! {panel_choice_label} में लॉगिन हो रहे हैं...")
                        st.balloons()
                        st.rerun()
                    else:
                        st.error("❌ गलत पासवर्ड! कृपया सही पासवर्ड डालें।")

        st.markdown(
            "<p style='text-align:center; color:#a3adbd; font-size:12px; margin-top:10px;'>"
            "🔐 आपका डेटा सुरक्षित है — हर पैनल का अपना अलग पासवर्ड है</p>",
            unsafe_allow_html=True
        )
    st.stop()

# =========================================================================
# 🔄 लॉगिन किए गए पैनल को लोड करना
# =========================================================================
panel = st.session_state["panel"]
panel_key = st.session_state.get("panel_key")

st.sidebar.markdown(
    f"""
    <div style="background: linear-gradient(135deg, #eef6ff, #f3fff5); border: 1px solid #d7e6f9;
                border-radius: 12px; padding: 12px 14px; margin-bottom: 10px;">
        <div style="font-size:11px; color:#7a869f; font-weight:600; letter-spacing:0.5px;">लॉगिन पैनल</div>
        <div style="font-size:15px; font-weight:700; color:#1a3c6e; margin-top:2px;">{panel}</div>
    </div>
    """,
    unsafe_allow_html=True
)
if st.sidebar.button("🔓 Logout करें", use_container_width=True):
    st.session_state["ok"] = False
    st.session_state.pop("panel", None)
    st.session_state.pop("panel_key", None)
    st.session_state["show_farewell"] = True
    st.rerun()

# 👁️ सिर्फ Admin (P6 लॉगिन) के लिए: बाकी सभी पैनल्स (P1-P5) को भी देखने का विकल्प
is_admin_session = (panel_key == "p6")
if is_admin_session:
    st.sidebar.divider()
    admin_view_choice = st.sidebar.selectbox(
        "👁️ पैनल देखें (Admin View):",
        ["⚙️ 6. Admin Panel"] + [PANEL_KEY_TO_NAME[k] for k in ["p1", "p2", "p3", "p4", "p5"]],
        key="admin_view_selector"
    )
    active_panel = admin_view_choice
else:
    active_panel = panel

import io
from openpyxl.styles import PatternFill, Border, Side

def generate_colored_excel_bytes(df_filtered, deg_col, br_col, minor_col_found, mdc_col_found, voc_col_found, pw_col_found, master_rules=None, sheet_name="Verified_Data"):
    """
    🔧 रीयूज़ेबल फ़ंक्शन: किसी भी DataFrame को रंगीन (🔴 गलत / 🔵 खाली) Excel bytes में बदलता है।
    Panel 3/4 और Admin Panel — दोनों जगह इसी फ़ंक्शन का इस्तेमाल होता है, ताकि रंग-कोडिंग हमेशा एक जैसी रहे।
    """
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_filtered.to_excel(writer, index=False, sheet_name=sheet_name)
        workbook = writer.book
        worksheet = writer.sheets[sheet_name]

        blue_fill = PatternFill(start_color="D1ECF1", end_color="D1ECF1", fill_type="solid")  # ब्लैंक = नीला
        red_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")    # गलत = लाल
        thin_border = Border(left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'),
                             top=Side(style='thin', color='CCCCCC'), bottom=Side(style='thin', color='CCCCCC'))

        targets_xl = {minor_col_found: 'minor', mdc_col_found: 'mdc', voc_col_found: 'voc', pw_col_found: 'pw'}

        for idx, (_, row) in enumerate(df_filtered.iterrows()):
            row_num = idx + 2  # एक्सेल डेटा रो

            deg_part = str(row[deg_col]) if deg_col and deg_col in df_filtered.columns else ""
            br_part = str(row[br_col]) if br_col and br_col in df_filtered.columns else ""
            student_deg = (deg_part + " " + br_part).lower().replace(".", "").replace(" ", "").strip()

            matched_key = "Default"
            if master_rules:
                sorted_keys = sorted(master_rules.keys(), key=len, reverse=True)
                for rule_key in sorted_keys:
                    rule_words = [w.lower().replace(".", "").strip() for w in rule_key.split() if w.strip()]
                    if rule_words and all(w in student_deg for w in rule_words):
                        matched_key = rule_key
                        break
            c_rule = master_rules.get(matched_key, {"minor": [], "mdc": [], "voc": [], "pw": []}) if master_rules else {"minor": [], "mdc": [], "voc": [], "pw": []}

            for col_idx, col_name in enumerate(df_filtered.columns, start=1):
                cell = worksheet.cell(row=row_num, column=col_idx)
                val = row[col_name]

                if col_name == br_col:
                    if pd.isna(val) or str(val).strip() == "":
                        cell.fill = blue_fill
                        cell.border = thin_border
                    continue

                if col_name in targets_xl:
                    rule_key = targets_xl[col_name]

                    if pd.isna(val) or str(val).strip() == "":
                        cell.fill = blue_fill
                        cell.border = thin_border
                    else:
                        val_clean = str(val).strip().lower().replace(".", "").replace(" ", "")
                        valid_list = c_rule.get(rule_key, [])
                        valid_set = {str(x).strip().lower().replace(".", "").replace(" ", "") for x in valid_list}

                        if valid_set and (val_clean not in valid_set):
                            cell.fill = red_fill
                            cell.border = thin_border

    return output.getvalue()

def process_panel_validation(df_panel, prefix, allowed_degrees, master_rules=None):
    deg_col = next((c for c in df_panel.columns if any(k in c.lower() for k in ['deg', 'course', 'class'])), df_panel.columns[0])
    br_col = next((c for c in df_panel.columns if any(k in c.lower() for k in ['branch', 'stream', 'subject'])), df_panel.columns[0])
    
    def check_degree(val):
        v = str(val).lower().replace(".", "").replace(" ", "").strip()
        return any(d in v for d in allowed_degrees)
        
    df_filtered = df_panel[df_panel[deg_col].apply(check_degree)].reset_index(drop=True)
    
    if df_filtered.empty:
        st.warning(f"⚠️ {prefix.upper()} पैनल के लिए कोई उपयुक्त डेटा (मैचिंग डिग्री) नहीं मिला।")
        return

    minor_col_found = next((c for c in df_filtered.columns if 'minor' in c.lower()), None)
    mdc_col_found = next((c for c in df_filtered.columns if 'mdc' in c.lower()), None)
    voc_col_found = next((c for c in df_filtered.columns if 'voc' in c.lower() or 'skill' in c.lower()), None)
    pw_col_found = next((c for c in df_filtered.columns if any(k in c.lower() for k in ['pw', 'project', 'ce'])), None)

    # --- 🖥️ लाइव वैरिफिकेशन स्टाइलर फ़ंक्शन (स्क्रीन ग्रिड के लिए फिक्स) ---
    def cell_styler(dataframe):
        s_df = pd.DataFrame('', index=dataframe.index, columns=dataframe.columns)
        targets = {minor_col_found: 'minor', mdc_col_found: 'mdc', voc_col_found: 'voc', pw_col_found: 'pw'}
        
        for index, row in dataframe.iterrows():
            # 🔧 फिक्स: Degree column + Branch column दोनों को मिलाकर चेक करना
            # (Biotechnology / Commerce Computer जैसी ब्रांच अक्सर अलग Branch column में होती है, Degree column में नहीं)
            deg_part = str(row[deg_col]) if deg_col else ""
            br_part = str(row[br_col]) if br_col else ""
            student_deg = (deg_part + " " + br_part).lower().replace(".", "").replace(" ", "").strip()
            
            matched_key = "Default"
            if master_rules:
                sorted_keys = sorted(master_rules.keys(), key=len, reverse=True)
                for rule_key in sorted_keys:
                    # 🔧 फिक्स: पूरा नाम एक साथ ढूंढने के बजाय हर word अलग-अलग ढूंढना
                    # (जैसे "B.Com. Computer" -> "bcom" और "computer" दोनों कहीं भी मिलने चाहिए)
                    rule_words = [w.lower().replace(".", "").strip() for w in rule_key.split() if w.strip()]
                    if rule_words and all(w in student_deg for w in rule_words):
                        matched_key = rule_key
                        break
            
            c_rule = master_rules.get(matched_key, {"minor": [], "mdc": [], "voc": [], "pw": []}) if master_rules else {"minor": [], "mdc": [], "voc": [], "pw": []}
            
            # ब्रांच की खाली चेकिंग
            if br_col and br_col in dataframe.columns:
                b_val = row[br_col]
                if pd.isna(b_val) or str(b_val).strip() == "":
                    s_df.at[index, br_col] = 'background-color: #d1ecf1; color: #0c5460; font-weight: bold; border: 1px solid #17a2b8;'
            
            # स्क्रीन पर नियमों के अनुसार सटीक लाइव कलर कोडिंग
            for col_name, rule_key in targets.items():
                if col_name and col_name in dataframe.columns:
                    val = row[col_name]
                    
                    # 🔵 स्थिति 1: अगर पूरी तरह से ब्लैंक है तो नीला करें
                    if pd.isna(val) or str(val).strip() == "":
                        s_df.at[index, col_name] = 'background-color: #d1ecf1; color: #0c5460; font-weight: bold; border: 1px solid #17a2b8;'
                    # 🔴 स्थिति 2: अगर भरा हुआ विषय नियमों से बाहर है तो लाल करें
                    else:
                        val_clean = str(val).strip().lower().replace(".", "").replace(" ", "")
                        valid_list = c_rule.get(rule_key, [])
                        valid_set = {str(x).strip().lower().replace(".", "").replace(" ", "") for x in valid_list}
                        if valid_set and (val_clean not in valid_set):
                            s_df.at[index, col_name] = 'background-color: #f8d7da; color: #721c24; font-weight: bold; border: 2px solid red;'
        return s_df

    st.subheader(f"📊 लाइव वैरिफाइड {prefix.upper()} डेटा टेबल")
    st.caption("🔵 नीला सेल = डेटा गायब है | 🔴 लाल सेल = गलत विषय (मास्टर गाइडलाइन से मिसमैच)")
    
    # स्क्रीन पर सीरियल नंबर 1 से शुरू करना
    df_filtered.index = range(1, len(df_filtered) + 1)
    st.dataframe(df_filtered.style.apply(cell_styler, axis=None), height=500, use_container_width=True)

    # --- 🚨 📥 रंगीन एक्सेल डाउनलोड (अब शेयर्ड फ़ंक्शन का इस्तेमाल कर रहा है) 🚨 ---
    processed_data = generate_colored_excel_bytes(
        df_filtered, deg_col, br_col, minor_col_found, mdc_col_found, voc_col_found, pw_col_found,
        master_rules=master_rules, sheet_name="Verified_Data"
    )
    st.download_button(
        label=f"📥 रंगीन (🔴/🔵) {prefix.upper()} डेटा एक्सेल डाउनलोड करें",
        data=processed_data,
        file_name=f"Verified_{prefix.upper()}_Colored_Data.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key=f"download_validated_excel_{prefix}"
    )

# =========================================================================
# 📥 PANEL 1: ENTRY / UPLOAD PANEL (डेटा सुरक्षित अपलोड)
# =========================================================================
if active_panel == "📥 1. Entry / Upload Panel":
    st.title("📥 Entry Panel - डेटा सुरक्षित अपलोड")
    if is_panel_hidden("p1") and not is_admin_session:
        st.warning("🔒 यह पैनल फिलहाल Admin द्वारा Hide किया गया है। डेटा उपलब्ध नहीं है।")
        st.dataframe(pd.DataFrame(), use_container_width=True)
        st.stop()
    st.write("यहाँ अपनी मुख्य एक्सेल/CSV फ़ाइल अपलोड करें। यह डेटा सीधे समीक्षा और क्लीनिंग के लिए **Work / Approve Panel (P2)** में ट्रांसफर हो जाएगा।")
    
    # एक्सेल या सीएसवी फ़ाइल अपलोड करने का विकल्प
    f = st.file_uploader("अपनी फ़ाइल अपलोड करें", type=["csv", "xlsx"])
    if f:
        try:
            # फ़ाइल टाइप के अनुसार डेटाबेस में रीड करना (Pandas Dataframe)
            df = pd.read_csv(f) if f.name.endswith('.csv') else pd.read_excel(f)
            st.success(f"🎉 फ़ाइल सफलतापूर्वक लोड हो गई ({len(df)} रोज़)!")
            
            # डेटा को P2 में ट्रांसफर करने का बटन
            if st.button("📤 वर्क/अप्रूवल पैनल (P2) में ट्रांसफर करें"):
                # पुराने किसी भी रॉ (Temporary) डेटा को साफ़ करना
                cursor.execute("DELETE FROM raw_store")
                
                # डेटाफ़्रेम को JSON में बदलकर सुरक्षित रूप से अस्थायी डेटाबेस में स्टोर करना
                cursor.execute("INSERT INTO raw_store (data_json) VALUES (?)", (json.dumps(df.to_dict(orient='records')),))
                conn.commit()
                
                # पुराने फ़ाइल के डिलीट किए गए कॉलम्स की सेटिंग्स को रीसेट करना
                st.session_state["deleted_cols"] = [] 
                
                st.success("🎉 डेटा सफलतापूर्वक अपलोड होकर **Work / Approve Panel** में प्रोसेस होने के लिए ट्रांसफर हो गया है!")
                st.balloons()
                st.rerun()
        except Exception as e:
            st.error(f"त्रुटि: {e}")

# =========================================================================
# 💻 PANEL 2: WORK / APPROVE PANEL (कॉलम मूव + लाइव स्प्लिट + डेटाबेस रूटिंग)
# =========================================================================
elif active_panel == "💻 2. Work / Approve Panel":
    st.title("💻 Work / Approve Panel - डेटा प्रोसेसिंग एवं अप्रूवल")
    if is_panel_hidden("p2") and not is_admin_session:
        st.warning("🔒 यह पैनल फिलहाल Admin द्वारा Hide किया गया है। डेटा उपलब्ध नहीं है।")
        st.dataframe(pd.DataFrame(), use_container_width=True)
        st.stop()
    
    # Panel 1 से ट्रांसफर होकर आया हुआ Staging (Raw) डेटा लोड करना
    raw_df = load_raw_data()
    
    if raw_df is None or raw_df.empty:
        st.info("📥 वर्तमान में कोई नई अपलोड की गई फ़ाइल पेंडिंग नहीं है। कृपया पहले 'Entry / Upload Panel (P1)' से फ़ाइल अपलोड करें।")
    else:
        # --- कार्य 1: बेकार कॉलम हटाना ---
        st.subheader("🗑️ बेकार कॉलम हटाएं (Remove Unwanted Columns)")
        active_cols = [c for c in raw_df.columns if c not in st.session_state["deleted_cols"]]
        
        cols_to_delete = st.multiselect("हटाने के लिए अनुपयोगी कॉलम चुनें:", options=active_cols)
        if cols_to_delete:
            if st.button("🔴 चुने गए कॉलम हटाएं"):
                st.session_state["deleted_cols"].extend(cols_to_delete)
                st.success("चयनित कॉलम स्क्रीन से हटा दिए गए!")
                st.rerun()
        
        final_raw_df = raw_df[active_cols]

        # --- कार्य 2: 🔄 कॉलमों का क्रम बदलना (Left/Right Move Feature) ---
        st.divider()
        st.subheader("🔄 कॉलमों का क्रम बदलें (Move Columns Left/Right)")
        st.write("नीचे दिए गए बॉक्स में क्रम बदलकर कॉलम को आगे-पीछे सेट करें। अप्रूवल के बाद इसी क्रम में लिस्ट लॉक होगी:")
        
        reordered_cols = st.multiselect(
            "कॉलमों का नया क्रम तय करें (सभी आवश्यक कॉलम इसी क्रम में चुनें):",
            options=active_cols,
            default=active_cols,
            key="col_reorder_select"
        )
        
        missing_cols = [c for c in active_cols if c not in reordered_cols]
        if missing_cols:
            reordered_cols.extend(missing_cols)
            
        final_raw_df = final_raw_df[reordered_cols]
        
        deg_col = next((c for c in final_raw_df.columns if any(k in c.lower() for k in ['deg', 'course', 'class'])), final_raw_df.columns)
        br_col = next((c for c in final_raw_df.columns if any(k in c.lower() for k in ['branch', 'stream', 'subject'])), final_raw_df.columns if len(final_raw_df.columns) > 1 else final_raw_df.columns)
        
        # --- कार्य 3: विशिष्ट डिग्री / ब्रांच की पूरी रो डिलीट करना ---
        st.divider()
        st.subheader("❌ विशिष्ट डिग्री / ब्रांच की पूरी रो डिलीट करें")
        st.write("यदि आप किसी खास कोर्स या स्ट्रीम का पूरा डेटा हटाना चाहते हैं, तो यहाँ से चुनें:")
        
        c_row1, c_row2 = st.columns(2)
        with c_row1:
            unique_degrees = final_raw_df[deg_col].dropna().unique().tolist()
            selected_degs = st.multiselect("डिलीट करने के लिए डिग्री (Course) चुनें:", options=unique_degrees)
        with c_row2:
            unique_branches = final_raw_df[br_col].dropna().unique().tolist()
            selected_branches = st.multiselect("डिलीट करने के लिए ब्रांच (Stream) चुनें:", options=unique_branches)
            
        if selected_degs or selected_branches:
            if st.button("🗑️ चुनी हुई रोज़ हमेशा के लिए डिलीट करें"):
                filtered_rows = []
                for _, row in raw_df.iterrows():
                    match_deg = str(row[deg_col]) in selected_degs if selected_degs else False
                    match_br = str(row[br_col]) in selected_branches if selected_branches else False
                    if not (match_deg or match_br):
                        filtered_rows.append(row.to_dict())
                
                cursor.execute("DELETE FROM raw_store")
                if filtered_rows:
                    cursor.execute("INSERT INTO raw_store (data_json) VALUES (?)", (json.dumps(filtered_rows),))
                conn.commit()
                st.success("🎉  चयनित डिग्री/ब्रांच की सभी रोज़ को सफलतापूर्वक डिलीट कर दिया गया है!")
                st.rerun()

        # फ़िल्टर्ड और रीऑर्डर किए गए डेटा का लाइव प्रीव्यू दिखाना
        st.divider()
        st.subheader("📋 अपलोड किए गए रॉ डेटा का लाइव प्रीव्यू (संशोधित क्रम)")
        st.dataframe(final_raw_df, height=350, use_container_width=True)
        
        el_col = next((c for c in final_raw_df.columns if any(k in c.lower() for k in ['elig', 'qual', 'class', 'course', 'deg'])), final_raw_df.columns)
        st.info(f"🔍 सिस्टम ऑटो-वर्गीकरण के लिए **'{el_col}'** कॉलम का उपयोग कर रहा है।")
        
        st.subheader("👀 लाइव प्री-विभाजन समीक्षा (Live Split Preview)")
        ug_preview_rows = []
        pg_preview_rows = []
        
        for _, row in final_raw_df.iterrows():
            val = str(row[el_col]).lower().replace(".", "").replace(" ", "").strip()
            if any(k in val for k in ["ma", "msc", "mcom", "mba", "mca", "post grad", "pg", "grad"]):
                pg_preview_rows.append(row.to_dict())
            else:
                ug_preview_rows.append(row.to_dict())
                
        df_ug_preview = pd.DataFrame(ug_preview_rows)
        df_pg_preview = pd.DataFrame(pg_preview_rows)
        
        prev_tab1, prev_tab2 = st.tabs([f"🎓 UG में जाने वाला डेटा ({len(df_ug_preview)} रोज़)", f"📜 PG में जाने वाला डेटा ({len(df_pg_preview)} रोज़)"])
        
        with prev_tab1:
            if not df_ug_preview.empty: st.dataframe(df_ug_preview, height=250, use_container_width=True)
            else: st.caption("कोई डेटा UG श्रेणी में नहीं मिला।")
        with prev_tab2:
            if not df_pg_preview.empty: st.dataframe(df_pg_preview, height=250, use_container_width=True)
            else: st.caption("कोई डेटा PG श्रेणी में नहीं मिला।")

        # --- कार्य 4: फाइनल अप्रूवल और रूटिंग एक्शन (नया क्रम डेटाबेस में लॉक होगा) ---
        st.divider()
        st.subheader("🚀 FINAL ACTION")
        st.write("📈 **डेटा ट्रांसफर:** क्लीन और रीऑर्डर किए गए छात्रों के डेटा को आगे UG (P3) और PG (P4) पैनल में भेजने के लिए यह बटन दबाएँ।")
        
        if st.button("✅ डेटा अप्रूव करें और पैनल्स में ट्रांसफर करें"):
            if not df_ug_preview.empty:
                ug_json_str = json.dumps(df_ug_preview[reordered_cols].to_dict(orient='records'))
                cursor.execute("INSERT INTO perma_store (data_json, course_type) VALUES (?, ?)", (ug_json_str, "UG"))
            if not df_pg_preview.empty:
                pg_json_str = json.dumps(df_pg_preview[reordered_cols].to_dict(orient='records'))
                cursor.execute("INSERT INTO perma_store (data_json, course_type) VALUES (?, ?)", (pg_json_str, "PG"))
            
            cursor.execute("DELETE FROM raw_store")
            conn.commit()
            st.success("🎉 बधाई हो! डेटा सफलतापूर्वक कस्टमाइज्ड क्रम में ट्रांसफर और लॉक कर दिया गया है।")
            st.balloons()
            st.rerun()

elif active_panel == "🎓 3. UG Panel":
    st.title("🎓 Undergraduate (UG) चेकिंग एवं त्रुटि सुधार पैनल")
    if is_panel_hidden("p3") and not is_admin_session:
        st.warning("🔒 यह पैनल फिलहाल Admin द्वारा Hide किया गया है। डेटा उपलब्ध नहीं है।")
        st.dataframe(pd.DataFrame(), use_container_width=True)
        st.stop()
    df_ug = load_permanent_data("UG")
    
    if df_ug is None or df_ug.empty: 
        st.info("ℹ️ UG डेटाबेस खाली है। कृपया पहले Panel 2 से डेटा अप्रूव करें।")
    else:
        # ऑटो-कॉलम डिटेक्शन
        minor_col = next((c for c in df_ug.columns if 'minor' in c.lower()), None)
        mdc_col = next((c for c in df_ug.columns if 'mdc' in c.lower()), None)
        voc_col = next((c for c in df_ug.columns if 'voc' in c.lower() or 'skill' in c.lower()), None)
        pw_col = next((c for c in df_ug.columns if any(k in c.lower() for k in ['pw', 'project', 'ce'])), None)
        
        # ड्रॉपडाउन में दिखाने के लिए यूनिक लिस्ट
        opt_minor = df_ug[minor_col].dropna().unique().tolist() if minor_col else []
        opt_mdc = df_ug[mdc_col].dropna().unique().tolist() if mdc_col else []
        opt_voc = df_ug[voc_col].dropna().unique().tolist() if voc_col else []
        opt_pw = df_ug[pw_col].dropna().unique().tolist() if pw_col else []
        
        st.markdown("### 🛠️ स्टेप 1: डिग्री-वाइज मास्टर गाइडलाइन सेट करें")
        st.caption("नीचे दी गई प्रत्येक डिग्री के बॉक्स को खोलकर उसके मान्य विषय चुनें और फिर 'लॉक करें' बटन दबाएं।")
        
        # --- डेटाबेस से पहले से सेव नियमों को सुरक्षित लोड करना ---
        ug_master_rules = {}
        try:
            cursor.execute("SELECT rules_json FROM locked_rules WHERE panel_prefix = 'ug_master'")
            locked_row = cursor.fetchone()
            if locked_row and locked_row[0]:
                ug_master_rules = json.loads(locked_row[0])
        except Exception as e:
            pass

        # आपकी मांगी गई 6 विशिष्ट डिग्रियां
        target_degrees = ["BA", "B.Sc.", "B.Sc. Biotechnology", "B.H.Sc.", "B.Com.", "B.Com. Computer"]
        current_configured_rules = {}
        
        for deg in target_degrees:
            with st.expander(f"📘 {deg} के लिए वैध विषय नियम (Valid Subjects)"):
                c1, c2, c3, c4 = st.columns(4)
                
                # पहले से सेव नियमों को ड्रॉपडाउन में डिफ़ॉल्ट दिखाना
                saved_deg_rule = ug_master_rules.get(deg, {})
                default_min = [x for x in saved_deg_rule.get("minor", []) if x in opt_minor]
                default_mdc = [x for x in saved_deg_rule.get("mdc", []) if x in opt_mdc]
                default_voc = [x for x in saved_deg_rule.get("voc", []) if x in opt_voc]
                default_pw = [x for x in saved_deg_rule.get("pw", []) if x in opt_pw]
                
                with c1:
                    r_minor = st.multiselect(f"Valid Minor", opt_minor, default=default_min, key=f"ug_min_{deg}")
                with c2:
                    r_mdc = st.multiselect(f"Valid MDC", opt_mdc, default=default_mdc, key=f"ug_mdc_{deg}")
                with c3:
                    r_voc = st.multiselect(f"Valid Vocational", opt_voc, default=default_voc, key=f"ug_voc_{deg}")
                with c4:
                    r_pw = st.multiselect(f"Valid PW/Ap/CE", opt_pw, default=default_pw, key=f"ug_pw_{deg}")
                    
                current_configured_rules[deg] = {
                    "minor": r_minor,
                    "mdc": r_mdc,
                    "voc": r_voc,
                    "pw": r_pw
                }
        
        # नियमों को डेटाबेस में लॉक करने का बटन
        if st.button("🔒 UG मास्टर विषय नियमावली लॉक करें", key="lock_master_ug_btn"):
            try:
                cursor.execute("""
                    INSERT INTO locked_rules (panel_prefix, rules_json) 
                    VALUES (?, ?)
                    ON CONFLICT(panel_prefix) DO UPDATE SET rules_json = excluded.rules_json
                """, ("ug_master", json.dumps(current_configured_rules)))
                conn.commit()
                st.success("🎉 सभी डिग्रियों के नियम डेटाबेस में सुरक्षित हो गए हैं और नीचे की लिस्ट रंगीन हो गई है!")
                st.rerun()
            except Exception as e:
                st.error(f"त्रुटि: {e}")
            
        st.divider()
        
        # लाइव वैलिडेशन टेबल रन करना (डेटाबेस से लोड किए गए नियमों को प्राथमिकता दें)
        rules_to_apply = ug_master_rules if ug_master_rules else current_configured_rules
        allowed_ug = ["ba", "bsc", "bcom", "bhsc", "bba", "bca", "computer"]
        
        process_panel_validation(df_ug, "ug", allowed_ug, master_rules=rules_to_apply)

# =========================================================================
# 📜 PANEL 4: PG PANEL
# =========================================================================
elif active_panel == "📜 4. PG Panel":
    st.title("📜 Postgraduate (PG) चेकिंग एवं त्रुटि सुधार पैनल")
    if is_panel_hidden("p4") and not is_admin_session:
        st.warning("🔒 यह पैनल फिलहाल Admin द्वारा Hide किया गया है। डेटा उपलब्ध नहीं है।")
        st.dataframe(pd.DataFrame(), use_container_width=True)
        st.stop()
    df_pg = load_permanent_data("PG")
    if df_pg is None or df_pg.empty: 
        st.info("ℹ️ PG डेटाबेस खाली है।")
    else: 
        process_panel_validation(df_pg, "pg", ["ma", "msc", "mcom", "mba", "mca", "post grad", "pg", "mtech", "llm"])

# =========================================================================
# 📊 PANEL 5: DASHBOARD / COUNTER PANEL (फुल स्क्रीन व्यूअर - भाग 1 और भाग 2 आपस में हाइड/शो)
# =========================================================================
elif active_panel == "📊 5. Dashboard / Counter Panel":
    # शीर्षक का फ़ॉन्ट छोटा किया गया है
    st.markdown("### 📊 Dashboard - डिग्री-वाइज लाइव काउंटर एवं विस्तृत डेटा समीक्षा")
    if is_panel_hidden("p5") and not is_admin_session:
        st.warning("🔒 यह पैनल फिलहाल Admin द्वारा Hide किया गया है। डेटा उपलब्ध नहीं है।")
        st.dataframe(pd.DataFrame(), use_container_width=True)
        st.stop()
    st.write("नीचे दिए गए टैब पर क्लिक करें, फिर बटन चुनकर 'विषय समरी' या 'छात्रों की फुल लिस्ट' को पूरी स्क्रीन पर देखें।")
    
    # 🛠️ फिक्स: UG और PG दोनों का डेटा लोड करना और लिस्ट बनाना
    df_ug_all = load_permanent_data("UG")
    df_pg_all = load_permanent_data("PG")
    
    # वेरिएबल को सही ढंग से इनिशियलाइज़ करना (यह डिलीट होने से एरर आ रहा था)
    all_dfs = []
    if df_ug_all is not None and not df_ug_all.empty: 
        all_dfs.append(df_ug_all)
    if df_pg_all is not None and not df_pg_all.empty: 
        all_dfs.append(df_pg_all)
    
    # अब यह कंडीशन बिना किसी एरर के बिल्कुल सही रन होगी
    if not all_dfs:
        st.info("ℹ️ काउंट प्रदर्शित करने के लिए डेटाबेस में कोई डेटा उपलब्ध नहीं है। कृपया पहले Panel 2 से डेटा अप्रूव करें।")
    else:
        master_df = pd.concat(all_dfs, ignore_index=True)
        
        # ऑटो-कॉलम डिटेक्शन इंजन
        deg_col = next((c for c in master_df.columns if any(k in c.lower() for k in ['deg', 'course', 'class'])), None)
        br_col = next((c for c in master_df.columns if any(k in c.lower() for k in ['branch', 'stream', 'subject'])), None)
        minor_col = next((c for c in master_df.columns if 'minor' in c.lower()), None)
        mdc_col = next((c for c in master_df.columns if 'mdc' in c.lower()), None)
        voc_col = next((c for c in master_df.columns if 'voc' in c.lower() or 'skill' in c.lower()), None)
        pw_col = next((c for c in master_df.columns if any(k in c.lower() for k in ['pw', 'project', 'ce'])), None)
        
        # 🔒 डेटाबेस से UG मास्टर नियमों को सुरक्षित लोड करना
        ug_master_rules = {}
        try:
            cursor.execute("SELECT rules_json FROM locked_rules WHERE panel_prefix = 'ug_master'")
            locked_row = cursor.fetchone()
            if locked_row and locked_row[0]:
                ug_master_rules = json.loads(locked_row[0])
        except:
            pass

        # डिग्रियों की सूची की सटीक मैपिंग (P3 - UG Rules Panel जैसी ही 6 डिग्री/ब्रांच संरचना)
        target_degrees = [
            {"display": "BA", "keywords": ["ba"]},
            {"display": "B.Sc.", "keywords": ["bsc"], "exclude": ["biotech"]},
            {"display": "B.Sc. Biotechnology", "keywords": ["bsc", "biotech"]},
            {"display": "B.H.Sc.", "keywords": ["bhsc"]},
            {"display": "B.Com.", "keywords": ["bcom"], "exclude": ["computer"]},
            {"display": "B.Com. Computer", "keywords": ["bcom", "computer"]}
        ]

        # सभी डिग्रियों के लिए इंटरएक्टिव टैब्स
        tab_titles = [deg["display"] for deg in target_degrees]
        tabs = st.tabs(tab_titles)

        for index, deg_info in enumerate(target_degrees):
            with tabs[index]:
                st.markdown(f"## 🎓 {deg_info['display']} डैशबोर्ड बोर्ड")
                
                # छात्र सूची में से इस विशिष्ट डिग्री के छात्रों को फ़िल्टर करना
                if deg_col and deg_col in master_df.columns:
                    def match_degree(row):
                        # 🔧 फिक्स: Degree column + Branch column दोनों को मिलाकर चेक करना
                        # (Biotechnology / Commerce Computer जैसी ब्रांच अक्सर अलग Branch column में होती है)
                        deg_part = str(row[deg_col]) if deg_col else ""
                        br_part = str(row[br_col]) if br_col and br_col in master_df.columns else ""
                        v = (deg_part + " " + br_part).lower().replace(".", "").replace(" ", "").strip()
                        match = all(k in v for k in deg_info["keywords"])
                        if "exclude" in deg_info:
                            if any(ex in v for ex in deg_info["exclude"]):
                                match = False
                        return match
                    
                    df_deg_filtered = master_df[master_df.apply(match_degree, axis=1)].reset_index(drop=True)
                else:
                    df_deg_filtered = pd.DataFrame()

                if df_deg_filtered.empty:
                    st.warning(f"⚠️ डेटाबेस में `{deg_info['display']}` का कोई छात्र रिकॉर्ड नहीं मिला।")
                    continue
                
                # मास्टर नियम लोड करना
                deg_rule = ug_master_rules.get(deg_info['display'], {"minor":[], "mdc":[], "voc":[], "pw":[]})

                # ✨ जादुई टॉगल बटन: यह तय करेगा कि भाग 1 देखना है या भाग 2
                view_option = st.radio(
                    "देखने के लिए व्यू चुनें:",
                    options=["📈 भाग 1: विषय काउंटर समरी (Subject Summary Counters)", "📋 भाग 2: छात्रों की विस्तृत लिस्ट (Detailed Student List)"],
                    horizontal=True,
                    key=f"view_toggle_{deg_info['display']}"
                )
                
                st.divider()

                # -------------------------------------------------------------------------
                # 📈 केवल भाग 1 (विषय काउंटर समरी) - लाइव कलर कोडिंग और अमान्य विषय रेड फिक्स
                # -------------------------------------------------------------------------
                if view_option == "📈 भाग 1: विषय काउंटर समरी (Subject Summary Counters)":
                    st.markdown("### 📈 विषयों की लाइव स्थिति (Summary Counters)")
                    
                    # समरी के अंदर माइनर, एमडीसी स्विच करने के लिए हॉरिजॉन्टल बार
                    selected_category = st.radio(
                        "समीक्षा के लिए विषय प्रकार चुनें:",
                        options=["Minor (माइनर)", "MDC (एम.डी.सी.)", "Vocational (व्यवसायिक)", "Project/PW (परियोजना)"],
                        horizontal=True,
                        key=f"cat_selector_{deg_info['display']}"
                    )
                
                    # 🔒 डेटाबेस से P3 (UG Panel) के नियमों को बिल्कुल सही तरीके से लोड करना
                    current_deg_rules = {}
                    try:
                        cursor.execute("SELECT rules_json FROM locked_rules WHERE panel_prefix = 'ug_master'")
                        locked_row = cursor.fetchone()
                        if locked_row and locked_row[0]:
                            all_rules = json.loads(locked_row[0])
                            current_deg_rules = all_rules.get(deg_info['display'], {"minor":[], "mdc":[], "voc":[], "pw":[]})
                    except Exception as e:
                        pass
                
                    cat_mapping = {
                        "Minor (माइनर)": {"col_name": minor_col, "rule_key": "minor", "label": "विषय का नाम (Minor Subject)"},
                        "MDC (एम.डी.सी.)": {"col_name": mdc_col, "rule_key": "mdc", "label": "विषय का नाम (MDC Subject)"},
                        "Vocational (व्यवसायिक)": {"col_name": voc_col, "rule_key": "voc", "label": "विषय का नाम (Vocational Subject)"},
                        "Project/PW (परियोजना)": {"col_name": pw_col, "rule_key": "pw", "label": "प्रोजेक्ट प्रकार (Project Type)"}
                    }
                
                    current_cat = cat_mapping[selected_category]
                
                    if current_cat["col_name"] and current_cat["col_name"] in df_deg_filtered.columns:
                        # काउंट्स (फ्रीक्वेंसी) की लाइव गणना
                        counts = df_deg_filtered[current_cat["col_name"]].dropna().value_counts().reset_index()
                        counts.columns = ["Subject", "Count"]
                        
                        if not counts.empty:
                            # P3 के लॉक नियमों से वैध विषयों का क्लीन सेट बनाना
                            valid_list = current_deg_rules.get(current_cat["rule_key"], [])
                            valid_set = {str(x).strip().lower().replace(".", "").replace(" ", "") for x in valid_list}
                            
                            # ✨ भाग 1 के लिए नया सख्त रो स्टाइलर इंजन (गलत विषय = चमकदार गाढ़ा लाल)
                            def row_styler(row):
                                sub_val = str(row["Subject"]).strip().lower().replace(".", "").replace(" ", "")
                                # यदि P3 में विषय चुने गए हैं और छात्र का विषय उसमें नहीं है, तो पूरी रो लाल होगी
                                if valid_set and (sub_val not in valid_set):
                                    return ['background-color: #f8d7da; color: #721c24; font-weight: bold; border: 2px solid #dc3545;'] * len(row)
                                return [''] * len(row)
                            
                            if current_cat["rule_key"] == "pw":
                                st.markdown("**Project Type Summary**<br><span style='color:gray; font-size:12px;'>(प्रोजेक्ट प्रकार की समरी सूची)</span>", unsafe_allow_html=True)
                            else:
                                st.markdown("**Subject Distribution Summary**<br><span style='color:gray; font-size:12px;'>(विषय आवंटन की समरी सूची)</span>", unsafe_allow_html=True)
                            
                            # फुल स्क्रीन चौड़ाई (Width) के साथ काउंटर तालिका रेंडर करना
                            # 🔧 फिक्स: टेबल की ऊंचाई अब रोज़ की संख्या के हिसाब से खुद-ब-खुद सेट होगी,
                            # ताकि पूरी लिस्ट एक बार में दिखे और स्क्रॉल न करना पड़े
                            dynamic_height = min(38 * (len(counts) + 1) + 3, 2000)
                            st.dataframe(
                                counts.style.apply(row_styler, axis=1), 
                                hide_index=True, 
                                use_container_width=True,
                                height=dynamic_height,
                                column_config={
                                    "Subject": st.column_config.TextColumn(label=current_cat["label"], width=600), 
                                    "Count": st.column_config.NumberColumn(label="छात्रों की संख्या (Count)", width=150)
                                }
                            )
                        else:
                            st.caption("इस श्रेणी में कोई डेटा उपलब्ध नहीं है।")
                    else:
                        st.caption("डेटाबेस में संबंधित कॉलम नहीं मिला।")

                # -------------------------------------------------------------------------
                # 📋 केवल भाग 2 (छात्रों की विस्तृत लिस्ट) - 1 से शुरू होने वाला सीरियल नंबर फिक्स
                # -------------------------------------------------------------------------
                elif view_option == "📋 भाग 2: छात्रों की विस्तृत लिस्ट (Detailed Student List)":
                    st.markdown(f"### 📋 {deg_info['display']} के सभी छात्रों का विस्तृत डेटा")
                    
                    df_to_show = df_deg_filtered.copy()
                    
                    # 🔍 लाइव सर्च बार फीचर
                    search_query = st.text_input(
                        f"🔍 {deg_info['display']} डेटा में सर्च करें (नाम, रोल नंबर या विषय डालें):", 
                        key=f"search_{deg_info['display']}"
                    )
                    
                    if search_query:
                        mask = df_to_show.astype(str).apply(lambda row: row.str.contains(search_query, case=False).any(), axis=1)
                        df_to_show = df_to_show[mask]
                
                    # --- 🛠️ सटीक लाइव कॉलम डिटेक्शन ---
                    actual_minor = next((c for c in df_to_show.columns if 'minor' in c.lower()), None)
                    actual_mdc = next((c for c in df_to_show.columns if 'mdc' in c.lower()), None)
                    actual_voc = next((c for c in df_to_show.columns if 'voc' in c.lower() or 'skill' in c.lower()), None)
                    actual_pw = next((c for c in df_to_show.columns if any(k in c.lower() for k in ['pw', 'project', 'ce'])), None)
                    actual_br = next((c for c in df_to_show.columns if any(k in c.lower() for k in ['branch', 'stream', 'subject'])), None)
                
                    # 🔒 डेटाबेस से P3 के नियमों को लोड करना
                    current_deg_rules = {}
                    try:
                        cursor.execute("SELECT rules_json FROM locked_rules WHERE panel_prefix = 'ug_master'")
                        locked_row = cursor.fetchone()
                        if locked_row and locked_row[0]:
                            all_rules = json.loads(locked_row[0])
                            current_deg_rules = all_rules.get(deg_info['display'], {"minor":[], "mdc":[], "voc":[], "pw":[]})
                    except Exception as e:
                        pass
                
                    # --- 📊 लाइव काउंटर मीटर ---
                    blank_count = 0
                    wrong_count = 0
                    
                    targets_for_counting = {
                        actual_minor: 'minor', 
                        actual_mdc: 'mdc', 
                        actual_voc: 'voc', 
                        actual_pw: 'pw'
                    }
                    
                    for index, row in df_to_show.iterrows():
                        for col_name, rule_key in targets_for_counting.items():
                            if col_name and col_name in df_to_show.columns:
                                val = row[col_name]
                                if pd.isna(val) or str(val).strip() == "":
                                    blank_count += 1
                                else:
                                    val_clean = str(val).strip().lower().replace(".", "").replace(" ", "")
                                    valid_list = current_deg_rules.get(rule_key, [])
                                    valid_set = {str(x).strip().lower().replace(".", "").replace(" ", "") for x in valid_list}
                                    if valid_set and (val_clean not in valid_set):
                                        wrong_count += 1
                
                    # स्क्रीन पर लाइव स्टेट्स कार्ड्स दिखाना
                    metric_c1, metric_c2, metric_c3 = st.columns(3)
                    with metric_c1:
                        st.metric(label="👥 कुल छात्र रिकॉर्ड (Total Rows)", value=len(df_to_show))
                    with metric_c2:
                        st.metric(label="🔵 कुल खाली सेल (Missing Data)", value=blank_count)
                    with metric_c3:
                        st.metric(label="🔴 कुल गलत विषय (Rule Mismatch)", value=wrong_count)
                
                    st.caption("🔵 नीला सेल = डेटा गायब है | 🔴 लाल सेल = गलत विषय (Panel 3 में आपके द्वारा चुने गए विषयों के अलावा बाकी सब)")
                
                    # ✨ जादू यहाँ है: टेबल दिखाने से पहले इंडेक्स को 1 से शुरू करने के लिए शिफ्ट करना
                    df_to_show.index = range(1, len(df_to_show) + 1)
                
                    # --- 🎨 लाइव कलर कोडिंग स्टाइलर इंजन ---
                    def full_table_styler(dataframe):
                        s_df = pd.DataFrame('', index=dataframe.index, columns=dataframe.columns)
                        targets = {
                            actual_minor: 'minor', 
                            actual_mdc: 'mdc', 
                            actual_voc: 'voc', 
                            actual_pw: 'pw'
                        }
                        for index, row in dataframe.iterrows():
                            row_has_wrong = False
                            for col_name, rule_key in targets.items():
                                if col_name and col_name in dataframe.columns:
                                    val = row[col_name]
                                    if pd.isna(val) or str(val).strip() == "":
                                        s_df.at[index, col_name] = 'background-color: #d1ecf1; color: #0c5460; font-weight: bold; border: 2px solid #17a2b8;'
                                    else:
                                        val_clean = str(val).strip().lower().replace(".", "").replace(" ", "")
                                        valid_list = current_deg_rules.get(rule_key, [])
                                        valid_set = {str(x).strip().lower().replace(".", "").replace(" ", "") for x in valid_list}
                                        if valid_set and (val_clean not in valid_set):
                                            s_df.at[index, col_name] = 'background-color: #f8d7da; color: #721c24; font-weight: bold; border: 2px solid #dc3545;'
                                            row_has_wrong = True
                            # 🟡 फिक्स: अगर इस रो में कोई भी सेल (Minor/MDC/Voc/PW) लाल (गलत) है,
                            # तो उसी रो के Branch सेल को पीला (Yellow) कर देना
                            if row_has_wrong and actual_br and actual_br in dataframe.columns:
                                s_df.at[index, actual_br] = 'background-color: #fff3cd; color: #856404; font-weight: bold; border: 2px solid #ffc107;'
                        return s_df
                
                    # full screen view table render
                    st.dataframe(
                        df_to_show.style.apply(full_table_styler, axis=None),
                        height=550,
                        use_container_width=True
                    )

                    # -------------------------------------------------------------------------
                    # 🟡 ब्रांच-वाइज सही/गलत समरी (जिस ब्रांच में कोई गलत छात्र है वो पीली दिखेगी)
                    # -------------------------------------------------------------------------
                    if actual_br and actual_br in df_to_show.columns:
                        st.divider()
                        st.markdown("### 🟡 ब्रांच-वाइज सही / गलत समरी")
                        st.caption("हर ब्रांच के सामने कुल कितने छात्र हैं, उनमें से कितने सही (✅) हैं और कितने गलत (🔴) हैं — जिस ब्रांच में कम-से-कम एक गलत छात्र है, वो रो पीली (Yellow) दिखेगी।")

                        targets_for_branch_summary = {
                            actual_minor: 'minor',
                            actual_mdc: 'mdc',
                            actual_voc: 'voc',
                            actual_pw: 'pw'
                        }

                        def _row_has_wrong_subject(row):
                            for col_name, rule_key in targets_for_branch_summary.items():
                                if col_name and col_name in df_to_show.columns:
                                    val = row[col_name]
                                    if not (pd.isna(val) or str(val).strip() == ""):
                                        val_clean = str(val).strip().lower().replace(".", "").replace(" ", "")
                                        valid_list = current_deg_rules.get(rule_key, [])
                                        valid_set = {str(x).strip().lower().replace(".", "").replace(" ", "") for x in valid_list}
                                        if valid_set and (val_clean not in valid_set):
                                            return True
                            return False

                        branch_summary_rows = []
                        for branch_val, group in df_to_show.groupby(df_to_show[actual_br].fillna("(खाली/Blank)").replace("", "(खाली/Blank)")):
                            total_n = len(group)
                            wrong_n = int(group.apply(_row_has_wrong_subject, axis=1).sum())
                            correct_n = total_n - wrong_n
                            branch_summary_rows.append({
                                "Branch (ब्रांच)": branch_val,
                                "कुल छात्र (Total)": total_n,
                                "✅ सही (Correct)": correct_n,
                                "🔴 गलत (Wrong)": wrong_n
                            })

                        branch_summary_df = pd.DataFrame(branch_summary_rows).sort_values(
                            "कुल छात्र (Total)", ascending=False
                        ).reset_index(drop=True)

                        def branch_summary_row_styler(row):
                            if row["🔴 गलत (Wrong)"] > 0:
                                return ['background-color: #fff3cd; color: #856404; font-weight: bold; border: 1px solid #ffc107;'] * len(row)
                            return ['background-color: #d4edda; color: #155724; font-weight: bold; border: 1px solid #28a745;'] * len(row)

                        branch_dyn_height = min(38 * (len(branch_summary_df) + 1) + 3, 1500)
                        st.dataframe(
                            branch_summary_df.style.apply(branch_summary_row_styler, axis=1),
                            hide_index=True,
                            use_container_width=True,
                            height=branch_dyn_height
                        )
                                    
# =========================================================================
# ⚙️ PANEL 6: ADMIN PANEL
# =========================================================================
elif active_panel == "⚙️ 6. Admin Panel":
    st.title("⚙️ Admin Panel - मास्टर डेटाबेस कंट्रोल")

    # =====================================================================
    # 🎨 लॉगिन स्क्रीन कस्टमाइज़ेशन (टाइटल + लोगो अपलोड)
    # =====================================================================
    st.subheader("🎨 लॉगिन स्क्रीन कस्टमाइज़ करें")
    st.caption("यहाँ से आप लॉगिन पेज पर दिखने वाला टाइटल बदल सकते हैं, और इमोजी की जगह अपना खुद का लोगो अपलोड कर सकते हैं।")

    logo_col, title_col = st.columns([1, 1.4])

    with logo_col:
        st.markdown("**🖼️ लोगो अपलोड करें**")
        _current_logo = get_app_setting("login_logo_b64")
        _current_mime = get_app_setting("login_logo_mime", "image/png")
        if _current_logo:
            st.image(f"data:{_current_mime};base64,{_current_logo}", caption="अभी का लोगो", width=260)
        else:
            st.info("अभी कोई लोगो नहीं है — डिफ़ॉल्ट इमोजी (🎓🔒) दिख रहा है।")

        uploaded_logo = st.file_uploader("नया लोगो चुनें (PNG/JPG)", type=["png", "jpg", "jpeg", "webp"], key="logo_uploader")
        lc1, lc2 = st.columns(2)
        with lc1:
            if st.button("💾 लोगो सेव करें", use_container_width=True, disabled=(uploaded_logo is None)):
                import base64 as _b64
                logo_bytes = uploaded_logo.getvalue()
                encoded = _b64.b64encode(logo_bytes).decode("utf-8")
                set_app_setting("login_logo_b64", encoded)
                set_app_setting("login_logo_mime", uploaded_logo.type or "image/png")
                st.success("🎉 लोगो सफलतापूर्वक सेव हो गया!")
                st.rerun()
        with lc2:
            if st.button("🗑️ लोगो हटाएं (इमोजी दिखाएं)", use_container_width=True, disabled=(not _current_logo)):
                delete_app_setting("login_logo_b64")
                delete_app_setting("login_logo_mime")
                st.success("लोगो हटा दिया गया, अब डिफ़ॉल्ट इमोजी दिखेगा।")
                st.rerun()

    with title_col:
        st.markdown("**✏️ टाइटल और सबटाइटल एडिट करें**")
        _cur_title = get_app_setting("login_title", "NEP Master Data System")
        _cur_subtitle = get_app_setting("login_subtitle", "अपना पैनल चुनें और आगे बढ़ने के लिए पासवर्ड डालें")
        new_title_input = st.text_input("लॉगिन पेज का टाइटल", value=_cur_title, key="login_title_input")
        new_subtitle_input = st.text_input("लॉगिन पेज का सबटाइटल", value=_cur_subtitle, key="login_subtitle_input")
        if st.button("💾 टाइटल सेव करें", key="save_login_title_btn"):
            set_app_setting("login_title", new_title_input.strip() or "NEP Master Data System")
            set_app_setting("login_subtitle", new_subtitle_input.strip() or "अपना पैनल चुनें और आगे बढ़ने के लिए पासवर्ड डालें")
            st.success("🎉 टाइटल सफलतापूर्वक अपडेट हो गया!")
            st.rerun()

    st.divider()

    # =====================================================================
    # 🔑 पैनल पासवर्ड अपडेट सिस्टम (6 पैनल्स)
    # =====================================================================
    st.subheader("🔑 पैनल पासवर्ड अपडेट करें")
    st.caption("यहाँ से आप किसी भी पैनल (P1-P6) का पासवर्ड बदल सकते हैं। बदलने के बाद उस पैनल में लॉगिन के लिए नया पासवर्ड इस्तेमाल होगा।")

    pw_cols = st.columns(3)
    new_pw_inputs = {}
    for i, (pk, pname) in enumerate(PANEL_KEY_TO_NAME.items()):
        with pw_cols[i % 3]:
            new_pw_inputs[pk] = st.text_input(f"{pname} का नया पासवर्ड", value="", type="password", key=f"pw_input_{pk}", placeholder="खाली छोड़ें तो नहीं बदलेगा")

    if st.button("🔐 पासवर्ड सेव करें", key="save_panel_passwords_btn"):
        updated_any = False
        for pk, new_pw in new_pw_inputs.items():
            if new_pw.strip():
                set_panel_password(pk, new_pw.strip())
                updated_any = True
        if updated_any:
            st.success("🎉 चुने गए पैनल्स के पासवर्ड सफलतापूर्वक अपडेट हो गए हैं!")
        else:
            st.info("कोई नया पासवर्ड नहीं डाला गया, कुछ भी नहीं बदला।")

    st.divider()

    # =====================================================================
    # 👁️ पैनल हाइड / अनहाइड सिस्टम (P1 से P5 तक)
    # =====================================================================
    st.subheader("👁️ पैनल Hide / Unhide करें (P1 से P5)")
    st.caption("जिस पैनल को Hide करेंगे, उसमें सही पासवर्ड डालने पर भी डेटा नहीं दिखेगा (सिर्फ पैनल का ढांचा दिखेगा)। Unhide करने पर डेटा फिर से दिखने लगेगा।")

    hide_cols = st.columns(5)
    hide_keys = ["p1", "p2", "p3", "p4", "p5"]
    new_hidden_state = {}
    for i, pk in enumerate(hide_keys):
        with hide_cols[i]:
            current_hidden = is_panel_hidden(pk)
            new_hidden_state[pk] = st.checkbox(f"🙈 {PANEL_KEY_TO_NAME[pk]} Hide करें", value=current_hidden, key=f"hide_chk_{pk}")

    if st.button("💾 Hide/Unhide सेटिंग सेव करें", key="save_hide_settings_btn"):
        for pk, hide_flag in new_hidden_state.items():
            set_panel_hidden(pk, hide_flag)
        st.success("🎉 Hide/Unhide सेटिंग सफलतापूर्वक सेव हो गई है!")
        st.rerun()

    st.divider()

    df_ug_download = load_permanent_data("UG")
    df_pg_download = load_permanent_data("PG")

    # 🔒 UG मास्टर रूल्स लोड करना ताकि Admin की रंगीन डाउनलोड भी बाकी पैनल्स जैसी सही हो
    _admin_ug_rules = {}
    try:
        cursor.execute("SELECT rules_json FROM locked_rules WHERE panel_prefix = 'ug_master'")
        _locked_row = cursor.fetchone()
        if _locked_row and _locked_row[0]:
            _admin_ug_rules = json.loads(_locked_row[0])
    except Exception:
        pass

    def _detect_cols(df):
        deg_c = next((c for c in df.columns if any(k in c.lower() for k in ['deg', 'course', 'class'])), df.columns[0] if len(df.columns) else None)
        br_c = next((c for c in df.columns if any(k in c.lower() for k in ['branch', 'stream', 'subject'])), None)
        min_c = next((c for c in df.columns if 'minor' in c.lower()), None)
        mdc_c = next((c for c in df.columns if 'mdc' in c.lower()), None)
        voc_c = next((c for c in df.columns if 'voc' in c.lower() or 'skill' in c.lower()), None)
        pw_c = next((c for c in df.columns if any(k in c.lower() for k in ['pw', 'project', 'ce'])), None)
        return deg_c, br_c, min_c, mdc_c, voc_c, pw_c

    st.subheader("📥 डेटाबेस बैकअप डाउनलोड करें")
    st.caption("🔵 नीला सेल = डेटा गायब है | 🔴 लाल सेल = गलत विषय (मास्टर गाइडलाइन से मिसमैच)")
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("#### 🎓 UG डेटा बैकअप")
        if df_ug_download is not None and not df_ug_download.empty:
            st.download_button(
                label="📥 UG डेटा CSV डाउनलोड करें (सादा)", 
                data=df_ug_download.to_csv(index=False).encode('utf-8'), 
                file_name="Approved_UG_Data_Backup.csv", 
                mime="text/csv",
                key="admin_ug_csv_dl"
            )
            deg_c, br_c, min_c, mdc_c, voc_c, pw_c = _detect_cols(df_ug_download)
            ug_colored = generate_colored_excel_bytes(
                df_ug_download, deg_c, br_c, min_c, mdc_c, voc_c, pw_c,
                master_rules=_admin_ug_rules, sheet_name="UG_Backup"
            )
            st.download_button(
                label="📥 रंगीन (🔴/🔵) UG डेटा एक्सेल डाउनलोड करें",
                data=ug_colored,
                file_name="Approved_UG_Colored_Backup.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="admin_ug_colored_dl"
            )
        else: 
            st.info("UG डेटाबेस खाली है।")
            
    with c2:
        st.markdown("#### 📜 PG डेटा बैकअप")
        if df_pg_download is not None and not df_pg_download.empty:
            st.download_button(
                label="📥 PG डेटा CSV डाउनलोड करें (सादा)", 
                data=df_pg_download.to_csv(index=False).encode('utf-8'), 
                file_name="Approved_PG_Data_Backup.csv", 
                mime="text/csv",
                key="admin_pg_csv_dl"
            )
            deg_c, br_c, min_c, mdc_c, voc_c, pw_c = _detect_cols(df_pg_download)
            pg_colored = generate_colored_excel_bytes(
                df_pg_download, deg_c, br_c, min_c, mdc_c, voc_c, pw_c,
                master_rules=None, sheet_name="PG_Backup"
            )
            st.download_button(
                label="📥 रंगीन (🔵) PG डेटा एक्सेल डाउनलोड करें",
                data=pg_colored,
                file_name="Approved_PG_Colored_Backup.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="admin_pg_colored_dl",
                help="PG के लिए फिलहाल कोई मास्टर सब्जेक्ट नियम सेट नहीं है, इसलिए सिर्फ खाली सेल नीले दिखेंगे।"
            )
        else: 
            st.info("PG डेटाबेस खाली है।")

    st.divider()
    st.subheader("🚨 डेंजर ज़ोन")
    confirm_reset = st.checkbox("मैं पूरे सिस्टम (रॉ + अप्रूव्ड दोनों डेटाबेस) को रीसेट करने की पुष्टि करता हूँ।")
    also_delete_rules = st.checkbox("⚠️ लॉक किए गए सब्जेक्ट नियम (Minor/MDC/Voc/PW रूल्स) भी डिलीट करें (सामान्यतः इसे टिक न करें)")
    if st.button("💥 ऑल डेटाबेस रीसेट करें"):
        if confirm_reset:
            cursor.execute("DELETE FROM raw_store")
            cursor.execute("DELETE FROM perma_store")
            # 🔧 फिक्स: locked_rules अब डिफ़ॉल्ट रूप से डिलीट नहीं होगा, ताकि लॉक किए गए सब्जेक्ट नियम
            # नई फ़ाइल अपलोड करने के बाद भी सुरक्षित बने रहें
            if also_delete_rules:
                cursor.execute("DELETE FROM locked_rules")
            conn.commit()
            st.session_state["deleted_cols"] = []
            if also_delete_rules:
                st.success("सिस्टम पूरी तरह से रीसेट हो गया है (डेटा + लॉक किए गए नियम दोनों हट गए)!")
            else:
                st.success("डेटा रीसेट हो गया है! लॉक किए गए सब्जेक्ट नियम सुरक्षित रखे गए हैं।")
            st.rerun()
        else: 
            st.error("कृपया पहले पुष्टि चेकबॉक्स पर टिक करें।")

# =========================================================================
# 🏁 फुटर (हर पेज के नीचे प्रोफेशनल क्रेडिट लाइन)
# =========================================================================
render_footer()
